"""
MiniCPM challan OCR matcher.

Prediction path:
    challan_image.jpg -> MiniCPM OCR -> stock_items_rows__1_.csv stock matching

Benchmark path:
    challan_updated.xlsx is loaded only to report how close the predicted
    stock descriptions are to Column D.

Run:
    py test_minicpm.py
"""

from __future__ import annotations

import base64
import os
import re
import time
from dataclasses import dataclass
from pathlib import Path

import pandas as pd
import requests
from openpyxl import load_workbook
from rapidfuzz import fuzz


BASE_DIR = Path(__file__).resolve().parent
IMAGE_PATH = BASE_DIR / "challan_image.jpg"
STOCK_CSV = BASE_DIR / "stock_items_rows__1_.csv"
BENCHMARK_XLSX = BASE_DIR / "challan_updated.xlsx"

MODEL = os.getenv("OLLAMA_MODEL", "openbmb/minicpm-v4.5")
OLLAMA_CHAT_URL = os.getenv("OLLAMA_CHAT_URL", "http://localhost:11434/api/chat")
REQUEST_TIMEOUT = int(os.getenv("REQUEST_TIMEOUT", "600"))

PROMPT = """Read this handwritten industrial tools challan.
Return ONLY a numbered list of line items.
Format each line exactly as: NUMBER. FULL_DESCRIPTION | QTY UNIT

Rules:
- These are industrial item names such as taps, dies, drills, deburring blades, countersink tools, HIKOKI tools, and tape items.
- Carefully preserve fractions like 9/16, metric sizes like 8x1.25, and brands like TOTEM, ADDISON, HIKOKI, ET.
- Expand ditto marks by repeating the full leading product family from the previous non-ditto line.
- Write one line per visible challan row.
- No explanation. No heading."""

BRAND_TOKENS = {
    "TOTEM",
    "ADDISON",
    "HIKOKI",
    "ET",
    "MIRANDA",
    "BOSCH",
    "YG",
    "STEELGRIP",
}
COMMON_TOKENS = {
    "HSS",
    "E",
    "TOOL",
    "TOOLS",
    "ITEM",
    "INDUSTRIAL",
    "THE",
    "OF",
    "A",
}
PRIORITY_TOKENS = {
    "TIN",
    "LH",
    "ROUND",
    "LONG",
    "BOT",
    "SET",
    "SEC",
    "PAIR",
    "SPPT",
    "SPFL",
    "TPR",
    "UNF",
    "UNC",
    "BSW",
    "BSF",
    "D371",
    "DIN371",
    "M35",
    "COUNTERSINK",
    "DEBURRING",
    "BLADE",
    "TAPE",
}
COLOR_TOKENS = ("BLUE", "RED", "GREEN", "YELLOW", "BLACK")
SPECIAL_CANDIDATE_TOKENS = {"LH", "TIN", "SPPT", "SPFL", "TPR", "BOT", "LONG"}
CANONICAL_TOKEN_MAP = {
    "DEBURING": "DEBURRING",
    "BLADES": "BLADE",
    "TAPS": "TAP",
    "DIES": "DIE",
    "DRILS": "DRILL",
}


@dataclass
class OCRItem:
    line_no: int
    raw_description: str
    qty_text: str
    description: str


@dataclass
class RankedMatch:
    name: str
    score: float


@dataclass
class PredictionRow:
    line_no: int
    raw_description: str
    normalized_description: str
    qty_text: str
    predicted_name: str
    expected_name: str
    status: str
    similarity: int
    confidence: int
    alternates: list[str]


def load_image_b64(path: Path) -> str:
    return base64.b64encode(path.read_bytes()).decode()


def collapse_spaces(value: str) -> str:
    return re.sub(r"\s+", " ", str(value).strip())


def normalize_text(value: str) -> str:
    text = str(value).upper()
    text = text.replace("H.SS", "HSS").replace("HSS-TAP", "HSS TAP")
    text = text.replace("HIS DIE", "HSS DIE").replace("HIS ", "HSS ")
    text = re.sub(r"[^A-Z0-9+/\.\-]+", " ", text)
    return collapse_spaces(text)


def canonicalize_tokens(tokens: list[str]) -> list[str]:
    return [CANONICAL_TOKEN_MAP.get(token, token) for token in tokens]


def tokenize(value: str) -> list[str]:
    return canonicalize_tokens(normalize_text(value).split())


def token_set(value: str) -> set[str]:
    return set(tokenize(value))


def extract_numbers(value: str) -> list[str]:
    return re.findall(r"\d+(?:/\d+)?(?:\.\d+)?", normalize_text(value))


def query_ollama_chat(image_b64: str) -> tuple[str, float]:
    payload = {
        "model": MODEL,
        "think": False,
        "messages": [
            {
                "role": "system",
                "content": "You are a strict OCR extraction engine. Return only the requested text.",
            },
            {
                "role": "user",
                "content": PROMPT,
                "images": [image_b64],
            },
        ],
        "stream": False,
        "options": {"temperature": 0},
    }

    start = time.time()
    resp = requests.post(OLLAMA_CHAT_URL, json=payload, timeout=REQUEST_TIMEOUT)
    elapsed = round(time.time() - start, 1)
    resp.raise_for_status()
    return resp.json()["message"]["content"], elapsed


def numbered_ocr_lines(raw_text: str) -> list[str]:
    return [line.strip() for line in raw_text.splitlines() if re.match(r"^\d+\.", line.strip())]


def split_desc_qty(line: str) -> tuple[str, str]:
    body = re.sub(r"^\d+\.\s*", "", line)
    left, _, right = body.partition("|")
    return collapse_spaces(left), collapse_spaces(right)


def is_ditto(value: str) -> bool:
    return bool(re.match(r'^\s*["\u201c\u201d]+', value))


def prefix_and_suffix(description: str) -> tuple[str, str]:
    tokens = tokenize(description)

    prefix_tokens = []
    for token in tokens:
        if re.search(r"\d", token):
            break
        prefix_tokens.append(token)

    suffix_tokens: list[str] = []
    for token in reversed(tokens):
        if re.search(r"\d", token):
            break
        if token in BRAND_TOKENS:
            suffix_tokens.insert(0, token)
        else:
            break

    return " ".join(prefix_tokens).strip(), " ".join(suffix_tokens).strip()


def apply_ocr_fixes(description: str) -> str:
    fixed = normalize_text(description)
    replacements = {
        "TOIEN": "TOTEM",
        "TOIEM": "TOTEM",
        "TOTAN": "TOTEM",
        "ADMISON": "ADDISON",
        "ADVISOR": "ADDISON",
        "HISEKDI": "HIKOKI",
        "HISEKI": "HIKOKI",
        "HISEKOI": "HIKOKI",
        "AGY ": "AG4 ",
        "AG 4": "AG-4",
        "AG4 ": "AG-4 ",
        "CHLDR": "DRILL",
        "CHLUDN": "DRILL",
        "CHLUD": "DRILL",
        "CHIKU": "DRILL",
        "IEX": "10X",
        "1EX": "10X",
        "SPP ": "SPPT ",
    }
    for src, dst in replacements.items():
        fixed = fixed.replace(src, dst)
    return collapse_spaces(fixed)


def parse_ocr_items(raw_text: str) -> list[OCRItem]:
    items: list[OCRItem] = []
    prev_prefix = ""
    prev_suffix = ""

    for line in numbered_ocr_lines(raw_text):
        match = re.match(r"^(\d+)\.", line)
        if not match:
            continue

        line_no = int(match.group(1))
        raw_desc, qty_text = split_desc_qty(line)
        desc_core = re.sub(r'^[\"\u201c\u201d\s]+', "", raw_desc).strip()

        if is_ditto(raw_desc) or re.match(r"^\d", desc_core):
            parts = [part for part in (prev_prefix, desc_core, prev_suffix) if part]
            raw_desc = " ".join(parts).strip()

        description = apply_ocr_fixes(raw_desc)
        prefix, suffix = prefix_and_suffix(description)
        if prefix:
            prev_prefix = prefix
        if suffix:
            prev_suffix = suffix

        items.append(
            OCRItem(
                line_no=line_no,
                raw_description=collapse_spaces(raw_desc),
                qty_text=normalize_text(qty_text),
                description=description,
            )
        )

    return items


def stock_unit_kind(unit: str, name: str) -> str:
    text = normalize_text(f"{unit} {name}")
    if "PAIR" in text:
        return "PAIR"
    if "SET" in text:
        return "SET"
    return "PIECE"


def load_stock(path: Path) -> pd.DataFrame:
    df = pd.read_csv(path)
    stock = df[["id", "name", "unit"]].dropna(subset=["name"]).copy()
    stock["name"] = stock["name"].astype(str)
    stock["unit"] = stock["unit"].fillna("").astype(str)
    stock["norm"] = stock["name"].map(normalize_text)
    stock["numbers"] = stock["norm"].map(extract_numbers)
    stock["tokens"] = stock["name"].map(tokenize)
    stock["token_set"] = stock["tokens"].map(set)
    stock["brand_tokens"] = stock["token_set"].map(lambda vals: vals & BRAND_TOKENS)
    stock["unit_kind"] = stock.apply(lambda row: stock_unit_kind(row["unit"], row["name"]), axis=1)
    return stock.reset_index(drop=True)


def load_benchmark(path: Path) -> list[str]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    corrected = []
    for row in ws.iter_rows(min_row=5, values_only=True):
        if not row[0]:
            continue
        corrected.append(collapse_spaces(row[3]))
    return corrected


def qty_kind(qty_text: str) -> str:
    qty = normalize_text(qty_text)
    if "PAIR" in qty:
        return "PAIR"
    if "SET" in qty or "SEP" in qty or "SEE" in qty:
        return "SET"
    return "PIECE"


def classify_family(query: str) -> str:
    q = normalize_text(query)
    if "STEELGRIP" in q or ("TAPE" in q and "3/4" in q):
        return "STEELGRIP"
    if "HIKOKI" in q or "AG-4" in q or "AG-7" in q:
        return "HIKOKI"
    if "COUNTERSINK" in q or q.startswith("C520") or q.startswith("CS20"):
        return "COUNTERSINK"
    if "DEBUR" in q or "BLADE" in q:
        return "DEBUR"
    if "DRILL" in q or "M35" in q:
        return "DRILL"
    if "DIE" in q:
        return "DIE"
    if "TAP" in q or "DIN37" in q or "DIN371" in q or "D371" in q:
        return "TAP"
    return "ALL"


def prepare_query(description: str, qty_text: str) -> str:
    query = normalize_text(description)

    if "DIN37" in query or "DIN371" in query:
        query = query.replace("DIN371", "D371").replace("DIN37", "D371")
        if "TAP" not in query:
            query = f"HSS-E TAP {query}"

    if "SPP" in query and "SPPT" not in query:
        query = query.replace("SPP", "SPPT")
    query = query.replace("SPPT-IN", "SPPT TIN")

    if "C520" in query or "CS20" in query:
        query = query.replace("C520", "CS20").replace("CS20", "CS20 COUNTERSINK TOOL")

    unit_mode = qty_kind(qty_text)
    if unit_mode == "SET":
        query = query.replace(" SEP ", " SET ").replace(" SEE ", " SET ")
    else:
        query = query.replace(" SEP ", " SEC ").replace(" SEE ", " SEC ")

    if unit_mode == "PIECE":
        query = re.sub(r"\bS\b$", "SEC", query)

    return collapse_spaces(query)


def build_family_views(stock: pd.DataFrame) -> dict[str, list[dict]]:
    recs = stock.to_dict("records")
    return {
        "ALL": recs,
        "TAP": [r for r in recs if "TAP" in r["token_set"]],
        "DIE": [r for r in recs if "DIE" in r["token_set"]],
        "DRILL": [r for r in recs if "DRILL" in r["token_set"]],
        "HIKOKI": [r for r in recs if "HIKOKI" in r["token_set"] or "AG-4" in r["token_set"] or "AG-7" in r["token_set"]],
        "COUNTERSINK": [r for r in recs if "COUNTERSINK" in r["token_set"]],
        "DEBUR": [r for r in recs if "DEBURRING" in r["token_set"] or "BLADE" in r["token_set"] or "DEBUR" in r["token_set"]],
        "STEELGRIP": [r for r in recs if "STEELGRIP" in r["token_set"] or "TAPE" in r["token_set"]],
    }


def meaningful_tokens(tokens: set[str]) -> set[str]:
    return {token for token in tokens if token not in COMMON_TOKENS and len(token) > 1}


def as_float(value: str) -> float | None:
    if "/" in value:
        return None
    try:
        return float(value)
    except ValueError:
        return None


def numeric_distance_penalty(query_numbers: list[str], candidate_numbers: list[str], family: str) -> float:
    if family not in {"TAP", "DIE", "DRILL", "COUNTERSINK"}:
        return 0.0

    query_values = [num for num in (as_float(value) for value in query_numbers) if num is not None]
    candidate_values = [num for num in (as_float(value) for value in candidate_numbers) if num is not None]
    if not query_values or not candidate_values:
        return 0.0

    nearest = min(abs(q - c) for q in query_values for c in candidate_values)
    return min(18.0, nearest * 10.0)


def family_keywords(family: str) -> set[str]:
    return {
        "TAP": {"TAP"},
        "DIE": {"DIE"},
        "DRILL": {"DRILL", "M35"},
        "HIKOKI": {"HIKOKI", "AG-4", "AG-7"},
        "COUNTERSINK": {"COUNTERSINK", "CS20"},
        "DEBUR": {"DEBURRING", "BLADE"},
        "STEELGRIP": {"STEELGRIP", "TAPE"},
    }.get(family, set())


def expand_steelgrip_colors(description: str) -> list[str]:
    normalized = normalize_text(description)
    tail = normalized.replace("STEELGRIP", " ")
    color_chunks = tail.split()
    code_map = {
        "B": ["BLUE"],
        "R": ["RED"],
        "G": ["GREEN"],
        "Y": ["YELLOW"],
        "BK": ["BLACK"],
        "BR": ["BLUE", "RED"],
        "GY": ["GREEN", "YELLOW"],
        "BY": ["BLUE", "YELLOW"],
    }

    colors: list[str] = []
    for chunk in color_chunks:
        letters = re.sub(r"[^A-Z]", "", chunk)
        if not letters or letters in {"STEELGRIP", "TAPE"}:
            continue
        if re.search(r"\d", chunk) and len(letters) <= 1:
            continue

        idx = 0
        while idx < len(letters):
            two = letters[idx : idx + 2]
            one = letters[idx]
            if two in code_map:
                colors.extend(code_map[two])
                idx += 2
                continue
            if one in code_map:
                colors.extend(code_map[one])
            idx += 1

    return colors


def candidate_pool(query: str, family: str, family_views: dict[str, list[dict]]) -> list[dict]:
    primary = family_views.get(family, family_views["ALL"])
    if family == "ALL":
        primary = family_views["ALL"]

    query_tokens = token_set(query)
    query_numbers = set(extract_numbers(query))
    query_brands = query_tokens & BRAND_TOKENS

    narrowed = []
    for record in primary:
        overlap_tokens = len(query_tokens & record["token_set"])
        overlap_numbers = len(query_numbers & set(record["numbers"]))
        overlap_brands = len(query_brands & record["brand_tokens"])
        if overlap_numbers or overlap_brands or overlap_tokens >= 2:
            narrowed.append(record)

    if len(narrowed) >= 12:
        return narrowed

    if family == "ALL":
        return primary

    seen_ids = {record["id"] for record in primary}
    fallback = list(primary)
    for record in family_views["ALL"]:
        if record["id"] in seen_ids:
            continue
        overlap_tokens = len(query_tokens & record["token_set"])
        overlap_numbers = len(query_numbers & set(record["numbers"]))
        overlap_brands = len(query_brands & record["brand_tokens"])
        if overlap_numbers or overlap_brands or overlap_tokens >= 3:
            fallback.append(record)
    return fallback


def score_candidate(query: str, qty_text: str, family: str, record: dict) -> float:
    candidate = record["norm"]
    query_tokens = token_set(query)
    candidate_tokens = record["token_set"]
    query_numbers = extract_numbers(query)
    candidate_numbers = record["numbers"]
    query_brands = query_tokens & BRAND_TOKENS
    candidate_brands = record["brand_tokens"]

    score = (
        0.55 * fuzz.WRatio(query, candidate)
        + 0.30 * fuzz.token_set_ratio(query, candidate)
        + 0.15 * fuzz.partial_ratio(query, candidate)
    )

    meaningful_overlap = meaningful_tokens(query_tokens) & meaningful_tokens(candidate_tokens)
    missing_meaningful = meaningful_tokens(query_tokens) - meaningful_tokens(candidate_tokens)
    score += len(meaningful_overlap) * 4.0
    score -= min(14.0, len(missing_meaningful) * 2.0)

    query_number_set = set(query_numbers)
    candidate_number_set = set(candidate_numbers)
    overlap_numbers = query_number_set & candidate_number_set
    missing_numbers = query_number_set - candidate_number_set
    extra_numbers = candidate_number_set - query_number_set
    score += len(overlap_numbers) * 16.0
    score -= len(missing_numbers) * 18.0
    if query_number_set:
        score -= min(10.0, len(extra_numbers) * 4.0)
        score -= numeric_distance_penalty(query_numbers, candidate_numbers, family)

    if query_brands:
        if query_brands & candidate_brands:
            score += 12.0
        else:
            score -= 10.0

    for token in PRIORITY_TOKENS:
        if token in query_tokens and token in candidate_tokens:
            score += 4.0
        elif token in query_tokens and token not in candidate_tokens:
            score -= 3.0

    unexpected_specials = {
        token for token in SPECIAL_CANDIDATE_TOKENS if token in candidate_tokens and token not in query_tokens
    }
    score -= len(unexpected_specials) * 4.0

    if family == "DIE" and "LH" not in query_tokens and "LH" in candidate_tokens:
        score -= 8.0
    if family == "DIE" and "ROUND" in candidate_tokens and "LH" not in query_tokens:
        score += 4.0

    unit_mode = qty_kind(qty_text)
    if unit_mode == record["unit_kind"]:
        score += 10.0
    elif unit_mode in {"SET", "PAIR"} and record["unit_kind"] == "PIECE":
        score -= 6.0

    hints = family_keywords(family)
    if hints:
        if hints & candidate_tokens:
            score += 10.0
        else:
            score -= 14.0

    return round(score, 2)


def confidence_from_matches(matches: list[RankedMatch]) -> int:
    if not matches:
        return 0

    top_score = matches[0].score
    second_score = matches[1].score if len(matches) > 1 else top_score - 12.0
    gap = max(0.0, top_score - second_score)
    confidence = 40.0 + min(30.0, max(0.0, top_score - 70.0) * 0.65) + min(25.0, gap * 1.6)
    return max(5, min(99, int(round(confidence))))


def rank_candidates(query: str, qty_text: str, family: str, family_views: dict[str, list[dict]], limit: int = 3) -> list[RankedMatch]:
    pool = candidate_pool(query, family, family_views)
    scored = [RankedMatch(record["name"], score_candidate(query, qty_text, family, record)) for record in pool]
    scored.sort(key=lambda match: match.score, reverse=True)
    return scored[:limit]


def rank_steelgrip_bundle(item: OCRItem, family_views: dict[str, list[dict]]) -> list[RankedMatch]:
    colors = expand_steelgrip_colors(item.description)
    if not colors:
        query = prepare_query(item.description, item.qty_text)
        return rank_candidates(query, item.qty_text, "STEELGRIP", family_views, limit=3)

    matches: list[RankedMatch] = []
    for color in colors:
        color_query = normalize_text(f'STEELGRIP TAPE 3/4" {color}')
        color_matches = rank_candidates(color_query, item.qty_text, "STEELGRIP", family_views, limit=1)
        if not color_matches:
            continue
        best = color_matches[0]
        matches.append(best)

    if not matches:
        query = prepare_query(item.description, item.qty_text)
        return rank_candidates(query, item.qty_text, "STEELGRIP", family_views, limit=3)

    joined_name = " + ".join(match.name for match in matches)
    average_score = round(sum(match.score for match in matches) / len(matches), 2)
    return [RankedMatch(joined_name, average_score)]


def predict_matches(item: OCRItem, family_views: dict[str, list[dict]]) -> list[RankedMatch]:
    family = classify_family(item.description)
    if family == "STEELGRIP":
        return rank_steelgrip_bundle(item, family_views)

    query = prepare_query(item.description, item.qty_text)
    return rank_candidates(query, item.qty_text, family, family_views, limit=3)


def compare_against_reference(predicted_name: str, expected_name: str) -> tuple[str, int]:
    if not expected_name:
        return "NO_REF", 0

    predicted_norm = normalize_text(predicted_name)
    expected_norm = normalize_text(expected_name)
    similarity = max(
        int(fuzz.WRatio(predicted_norm, expected_norm)),
        int(fuzz.token_set_ratio(predicted_norm, expected_norm)),
    )

    if predicted_norm == expected_norm:
        return "EXACT", 100

    predicted_tokens = token_set(predicted_name)
    expected_tokens = token_set(expected_name)
    predicted_numbers = set(extract_numbers(predicted_name))
    expected_numbers = set(extract_numbers(expected_name))
    shared_meaningful = meaningful_tokens(predicted_tokens) & meaningful_tokens(expected_tokens)
    predicted_family = classify_family(predicted_name)
    expected_family = classify_family(expected_name)
    same_family = predicted_family == expected_family
    brand_match = bool((predicted_tokens & BRAND_TOKENS) & (expected_tokens & BRAND_TOKENS))

    if similarity >= 92:
        return "CLOSE", similarity
    if same_family and predicted_numbers & expected_numbers and len(shared_meaningful) >= 2:
        return "CLOSE", similarity
    if same_family and brand_match and len(shared_meaningful) >= 2:
        return "CLOSE", similarity
    if predicted_family == "STEELGRIP" and expected_family == "STEELGRIP" and len(shared_meaningful & set(COLOR_TOKENS)) >= 2:
        return "CLOSE", similarity
    if same_family and len(shared_meaningful) >= 3:
        return "CLOSE", similarity

    return "DIFFERENT", similarity


def build_prediction_rows(items: list[OCRItem], expected: list[str], family_views: dict[str, list[dict]]) -> list[PredictionRow]:
    rows = []
    for idx, item in enumerate(items):
        matches = predict_matches(item, family_views)
        predicted_name = matches[0].name if matches else ""
        expected_name = expected[idx] if idx < len(expected) else ""
        status, similarity = compare_against_reference(predicted_name, expected_name)
        rows.append(
            PredictionRow(
                line_no=item.line_no,
                raw_description=item.raw_description,
                normalized_description=item.description,
                qty_text=item.qty_text,
                predicted_name=predicted_name,
                expected_name=expected_name,
                status=status,
                similarity=similarity,
                confidence=confidence_from_matches(matches),
                alternates=[match.name for match in matches[1:3]],
            )
        )
    return rows


def print_predictions(rows: list[PredictionRow], elapsed: float) -> None:
    total = len(rows)
    exact = sum(1 for row in rows if row.status == "EXACT")
    close = sum(1 for row in rows if row.status in {"EXACT", "CLOSE"})
    different = sum(1 for row in rows if row.status == "DIFFERENT")

    print("Loaded sample files.")
    print(f"Model        : {MODEL}")
    print(f"Rows found   : {total}")
    print(f"Exact        : {exact}/{total}")
    print(f"Close-or-better: {close}/{total}")
    print(f"Different    : {different}/{total}")
    print(f"Inference    : {elapsed}s")
    print("Prediction uses only challan_image.jpg + stock_items_rows__1_.csv.")
    print("Benchmark uses challan_updated.xlsx Column D only for review.")
    print("=" * 120)

    for row in rows:
        print(f"{row.line_no}.")
        print(f"  OCR       : {row.normalized_description}")
        print(f"  Qty       : {row.qty_text}")
        print(f"  Predicted : {row.predicted_name}")
        print(f"  Reference : {row.expected_name}")
        print(f"  Status    : {row.status} | similarity {row.similarity} | confidence {row.confidence}")
        if row.alternates and (row.status != "EXACT" or row.confidence < 75):
            print(f"  Options   : {' | '.join(row.alternates)}")
        print("-" * 120)


def main() -> None:
    print("Loading files...")
    image_b64 = load_image_b64(IMAGE_PATH)
    stock = load_stock(STOCK_CSV)
    expected = load_benchmark(BENCHMARK_XLSX)
    family_views = build_family_views(stock)

    raw_text, elapsed = query_ollama_chat(image_b64)
    items = parse_ocr_items(raw_text)
    rows = build_prediction_rows(items, expected, family_views)
    print_predictions(rows, elapsed)


if __name__ == "__main__":
    main()
